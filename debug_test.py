import sys, os, re, io
sys.path.insert(0, '.')
print("starting...", flush=True)
from docx import Document
from docx.shared import Pt, Inches, Cm, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import docx.parts.hdrftr
print("docx imported", flush=True)

doc = Document()
print("document created", flush=True)
style = doc.styles["Normal"]
print("got style", flush=True)
style.font.name = "Times New Roman"
style.font.size = Pt(11)
print("set font", flush=True)
pf = style.paragraph_format
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
print("set alignment", flush=True)

p = doc.add_paragraph()
print("added paragraph", flush=True)
p.add_run("Hello world")
print("added run", flush=True)

buf = io.BytesIO()
doc.save(buf)
print("saved to buffer", flush=True)
print("OK", flush=True)
