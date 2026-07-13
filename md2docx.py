#!/usr/bin/env python3
"""MD2DOCX — bidirectional Markdown ↔ DOCX converter."""

import sys
import os
import queue
import threading
import re
import io
import zipfile
import datetime
import glob as globmod
import copy

from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, Cm, Emu, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml
    import docx.parts.hdrftr
except ImportError:
    Document = None

try:
    from lxml import etree
except ImportError:
    etree = None

USE_GUI = False
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    USE_GUI = True
except ImportError:
    TkinterDnD = None
    tk = None


COMMENT_RE = re.compile(r"^<!--\s*COMMENT\s*\[(.+?)\]\s*:\s*(.+?)\s*-->$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
EXISTING_NUM_RE = re.compile(r"^\d+(?:\.\d+)*\.?\s+")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBERED_RE = re.compile(r"^(\s*)\d+\.\s+(.*)$")
CODE_FENCE_RE = re.compile(r"^```(\w*)$")
BLOCKQUOTE_RE = re.compile(r"^>\s?(.*)$")
TABLE_RE = re.compile(r"^\|.*\|$")
HR_RE = re.compile(r"^([-*_])\s*\1\s*\1[\s\1]*$")

INLINE_BOLD_ITALIC_RE = re.compile(r"\*\*\*(.+?)\*\*\*")
INLINE_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
INLINE_ITALIC_RE = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
INLINE_CODE_RE = re.compile(r"`([^`]+)`")
INLINE_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _clear_theme_fonts(style):
    rPr = style.element.find(qn("w:rPr"))
    if rPr is None:
        rPr = parse_xml(f'<w:rPr {nsdecls("w")}></w:rPr>')
        style.element.append(rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}></w:rFonts>')
        rPr.insert(0, rFonts)
    for attr in ["asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"]:
        try:
            del rFonts.attrib[qn(f"w:{attr}")]
        except KeyError:
            pass


def _set_run_font(run, name="Times New Roman", size=11, bold=False, italic=False, color=None, underline=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    if color:
        run.font.color.rgb = color
    rPr = run._element.find(qn("w:rPr"))
    if rPr is not None:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is not None:
            for attr in ["asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"]:
                try:
                    del rFonts.attrib[qn(f"w:{attr}")]
                except KeyError:
                    pass


def _add_inline_formatting(paragraph, text):
    tokens = []
    pos = 0
    while pos < len(text):
        m = INLINE_BOLD_ITALIC_RE.search(text, pos)
        bi_m = INLINE_BOLD_RE.search(text, pos)
        i_m = INLINE_ITALIC_RE.search(text, pos)
        c_m = INLINE_CODE_RE.search(text, pos)
        l_m = INLINE_LINK_RE.search(text, pos)
        matches = []
        if m:
            matches.append((m.start(), "bold_italic", m))
        if bi_m:
            matches.append((bi_m.start(), "bold", bi_m))
        if i_m:
            matches.append((i_m.start(), "italic", i_m))
        if c_m:
            matches.append((c_m.start(), "code", c_m))
        if l_m:
            matches.append((l_m.start(), "link", l_m))

        matches = [m for m in matches if m[0] >= pos]
        if not matches:
            tokens.append(("text", text[pos:]))
            break

        matches.sort(key=lambda x: x[0])
        best = matches[0]

        if best[0] > pos:
            tokens.append(("text", text[pos:best[0]]))

        _, kind, match = best
        if kind == "bold_italic":
            tokens.append(("bold_italic", match.group(1)))
            pos = match.end()
        elif kind == "bold":
            tokens.append(("bold", match.group(1)))
            pos = match.end()
        elif kind == "italic":
            tokens.append(("italic", match.group(1)))
            pos = match.end()
        elif kind == "code":
            tokens.append(("code", match.group(1)))
            pos = match.end()
        elif kind == "link":
            tokens.append(("link", match.group(1), match.group(2)))
            pos = match.end()
    for token in tokens:
        if token[0] == "text":
            run = paragraph.add_run(token[1])
            _set_run_font(run)
        elif token[0] == "bold_italic":
            run = paragraph.add_run(token[1])
            _set_run_font(run, bold=True, italic=True)
        elif token[0] == "bold":
            run = paragraph.add_run(token[1])
            _set_run_font(run, bold=True)
        elif token[0] == "italic":
            run = paragraph.add_run(token[1])
            _set_run_font(run, italic=True)
        elif token[0] == "code":
            run = paragraph.add_run(token[1])
            _set_run_font(run, name="Courier New", size=10)
        elif token[0] == "link":
            run = paragraph.add_run(token[1])
            _set_run_font(run, color=RGBColor(0x05, 0x63, 0xC1), underline=True)


def _configure_document(doc):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    pf = style.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _clear_theme_fonts(style)

    _create_style(doc, "Code Block", "Normal")
    _create_style(doc, "Document Title", "Normal")


def _create_style(doc, name, base_style):
    style = doc.styles.add_style(name, 1)
    style.base_style = doc.styles[base_style]
    return style


def _set_heading_font(style):
    style.font.name = "Times New Roman"
    _clear_theme_fonts(style)


def _apply_heading_fonts(doc):
    for i in range(1, 10):
        try:
            _set_heading_font(doc.styles[f"Heading {i}"])
        except KeyError:
            pass
    for sname in ["List Bullet", "List Number"]:
        try:
            sty = doc.styles[sname]
            sty.font.name = "Times New Roman"
            _clear_theme_fonts(sty)
        except KeyError:
            pass


def parse_md_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    _configure_document(doc)
    _apply_heading_fonts(doc)

    title_style = doc.styles["Document Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(20)
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(24)
    _clear_theme_fonts(title_style)

    code_style = doc.styles["Code Block"]
    code_style.font.name = "Courier New"
    code_style.font.size = Pt(9)
    code_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    i = 0
    n = len(lines)
    pending_comments = []
    para_count = 0
    title_done = False

    comment_list = []

    def add_paragraph(text, style=None):
        nonlocal para_count
        if style:
            p = doc.add_paragraph(style=style)
        else:
            p = doc.add_paragraph()
        if text:
            _add_inline_formatting(p, text)
        idx = para_count
        para_count += 1
        for pc_author, pc_text in pending_comments:
            comment_list.append((idx, pc_author, pc_text))
        pending_comments.clear()
        return p, idx

    while i < n:
        raw = lines[i]
        line = raw.rstrip("\n").rstrip("\r")

        cm = COMMENT_RE.match(line)
        if cm:
            pending_comments.append((cm.group(1), cm.group(2)))
            i += 1
            continue

        hm = HEADING_RE.match(line)
        if hm:
            level = len(hm.group(1))
            text = hm.group(2)
            text = EXISTING_NUM_RE.sub("", text, count=1)
            if not title_done and level == 1:
                p = doc.add_paragraph(style="Document Title")
                _add_inline_formatting(p, text)
                para_count += 1
                title_done = True
            else:
                eff_level = level - 1 if title_done else level
                eff_level = max(1, min(eff_level, 9))
                p = doc.add_paragraph(style=f"Heading {eff_level}")
                _add_inline_formatting(p, text)
                para_count += 1
            for pc_author, pc_text in pending_comments:
                comment_list.append((para_count - 1, pc_author, pc_text))
            pending_comments.clear()
            i += 1
            continue

        if BULLET_RE.match(line):
            bullet_lines = []
            while i < n:
                l = lines[i].rstrip("\n").rstrip("\r")
                bm = BULLET_RE.match(l)
                if not bm:
                    break
                bullet_lines.append(bm.group(2))
                i += 1
            for bline in bullet_lines:
                p = doc.add_paragraph(style="List Bullet")
                _add_inline_formatting(p, bline)
                para_count += 1
                for pc_author, pc_text in pending_comments:
                    comment_list.append((para_count - 1, pc_author, pc_text))
                pending_comments.clear()
            continue

        if NUMBERED_RE.match(line):
            num_lines = []
            while i < n:
                l = lines[i].rstrip("\n").rstrip("\r")
                nm = NUMBERED_RE.match(l)
                if not nm:
                    break
                num_lines.append(nm.group(2))
                i += 1
            for nline in num_lines:
                p = doc.add_paragraph(style="List Number")
                _add_inline_formatting(p, nline)
                para_count += 1
                for pc_author, pc_text in pending_comments:
                    comment_list.append((para_count - 1, pc_author, pc_text))
                pending_comments.clear()
            continue

        if CODE_FENCE_RE.match(line):
            code_lines = []
            i += 1
            while i < n and not CODE_FENCE_RE.match(lines[i]):
                code_lines.append(lines[i].rstrip("\n").rstrip("\r"))
                i += 1
            i += 1
            if code_lines:
                text = "\n".join(code_lines)
                p = doc.add_paragraph(style="Code Block")
                run = p.add_run(text)
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                para_count += 1
                for pc_author, pc_text in pending_comments:
                    comment_list.append((para_count - 1, pc_author, pc_text))
                pending_comments.clear()
            continue

        if BLOCKQUOTE_RE.match(line):
            bq_lines = []
            while i < n:
                l = lines[i].rstrip("\n").rstrip("\r")
                bqm = BLOCKQUOTE_RE.match(l)
                if not bqm:
                    break
                bq_lines.append(bqm.group(1))
                i += 1
            if bq_lines:
                text = " ".join(bq_lines)
                p = doc.add_paragraph()
                _add_inline_formatting(p, text)
                for run in p.runs:
                    run.font.italic = True
                    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                p.paragraph_format.left_indent = Inches(0.5)
                para_count += 1
                for pc_author, pc_text in pending_comments:
                    comment_list.append((para_count - 1, pc_author, pc_text))
                pending_comments.clear()
            continue

        if TABLE_RE.match(line) and line.count("|") >= 3:
            sep_re = re.compile(r"^\|[\s\-:|+]+\|$")
            table_rows = []
            while i < n:
                l = lines[i].rstrip("\n").rstrip("\r")
                if TABLE_RE.match(l) and l.count("|") >= 3:
                    if sep_re.match(l):
                        i += 1
                        continue
                    cells = [c.strip() for c in l.split("|")[1:-1]]
                    table_rows.append(cells)
                    i += 1
                else:
                    break
            if len(table_rows) >= 2:
                num_cols = max(len(r) for r in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
                tbl.style = "Table Grid"
                tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
                tbl.autofit = True
                tbl.columns[0].width = None
                tbl_element = tbl._tbl
                tblPr = tbl_element.find(qn("w:tblPr"))
                if tblPr is None:
                    tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}></w:tblPr>')
                    tbl_element.insert(0, tblPr)
                tblW = tblPr.find(qn("w:tblW"))
                if tblW is None:
                    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="5000" w:type="pct"/>')
                    tblPr.append(tblW)
                else:
                    tblW.set(qn("w:w"), "5000")
                    tblW.set(qn("w:type"), "pct")
                for row_idx, row_data in enumerate(table_rows):
                    for col_idx in range(num_cols):
                        cell = tbl.cell(row_idx, col_idx)
                        cell.text = ""
                        val = row_data[col_idx] if col_idx < len(row_data) else ""
                        p = cell.paragraphs[0]
                        _add_inline_formatting(p, val)
                        if row_idx == 0:
                            for run in p.runs:
                                run.font.bold = True
                tr_elements = tbl_element.findall(qn("w:tr"))
                if len(tr_elements) > 0:
                    first_tr = tr_elements[0]
                    trPr = first_tr.find(qn("w:trPr"))
                    if trPr is None:
                        trPr = parse_xml(f'<w:trPr {nsdecls("w")}></w:trPr>')
                        first_tr.insert(0, trPr)
                    tblHeader = trPr.find(qn("w:tblHeader"))
                    if tblHeader is None:
                        tblHeader = parse_xml(f'<w:tblHeader {nsdecls("w")}/>')
                        trPr.append(tblHeader)
                para_count += 1
                for pc_author, pc_text in pending_comments:
                    comment_list.append((para_count - 1, pc_author, pc_text))
                pending_comments.clear()
            continue

        if HR_RE.match(line):
            if i + 1 < n:
                next_line = lines[i + 1].strip()
                if HEADING_RE.match(next_line):
                    i += 1
                    continue
            p = doc.add_paragraph()
            pPr = p._element.find(qn("w:pPr"))
            if pPr is None:
                pPr = parse_xml(f'<w:pPr {nsdecls("w")}></w:pPr>')
                p._element.insert(0, pPr)
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="auto"/>'
                f'</w:pBdr>'
            )
            pPr.append(pBdr)
            para_count += 1
            for pc_author, pc_text in pending_comments:
                comment_list.append((para_count - 1, pc_author, pc_text))
            pending_comments.clear()
            i += 1
            continue

        if line.strip() == "":
            i += 1
            continue

        plain_lines = []
        while i < n:
            l = lines[i].rstrip("\n").rstrip("\r")
            if l.strip() == "":
                break
            if (HEADING_RE.match(l) or BULLET_RE.match(l) or NUMBERED_RE.match(l)
                    or CODE_FENCE_RE.match(l) or BLOCKQUOTE_RE.match(l)
                    or (TABLE_RE.match(l) and l.count("|") >= 3)
                    or HR_RE.match(l) or COMMENT_RE.match(l)):
                break
            plain_lines.append(l)
            i += 1

        if plain_lines:
            text = " ".join(plain_lines)
            add_paragraph(text)

    _add_footer(doc)

    buf = io.BytesIO()
    doc.save(buf)

    if comment_list:
        _inject_comments(buf, docx_path, comment_list)
    else:
        with open(docx_path, "wb") as f:
            f.write(buf.getvalue())

    _apply_replacements(docx_path)
    _inject_numbering(docx_path)


def _add_footer(doc):
    class CustomFooterPart(docx.parts.hdrftr.FooterPart):
        @classmethod
        def _default_footer_xml(cls):
            return (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
                ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
                '<w:p><w:pPr><w:jc w:val="center"/></w:pPr>'
                '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                '<w:r><w:instrText> PAGE </w:instrText></w:r>'
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                '<w:r><w:t>1</w:t></w:r>'
                '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
                '</w:p></w:ftr>'
            )

    docx.parts.hdrftr.FooterPart = CustomFooterPart

    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.clear()
    run = footer_para.add_run()
    run._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    run2 = footer_para.add_run()
    run2._element.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>'))
    run3 = footer_para.add_run()
    run3._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
    run4 = footer_para.add_run("1")
    run5 = footer_para.add_run()
    run5._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))


def _build_style_heading_map(raw_bytes):
    heading_map = {}
    try:
        root = etree.fromstring(raw_bytes)
    except Exception:
        return heading_map
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for style_el in root.iter(f"{{{ns}}}style"):
        style_id = style_el.get(f"{{{ns}}}styleId")
        name_el = style_el.find(f"{{{ns}}}name")
        if name_el is None:
            continue
        name_val = name_el.get(f"{{{ns}}}val", "")
        outline_el = style_el.find(f"{{{ns}}}pPr/{ns}outlineLvl" if False else f"{{{ns}}}pPr/{{{ns}}}outlineLvl")
        if outline_el is not None:
            level = outline_el.get(f"{{{ns}}}val")
            if level is not None:
                heading_map[style_id] = int(level) + 1
                continue
        lower_name = name_val.lower()
        for prefix in ("heading ", "заголовок "):
            if lower_name.startswith(prefix):
                try:
                    level = int(lower_name[len(prefix):])
                    if 1 <= level <= 9:
                        heading_map[style_id] = level
                except ValueError:
                    pass
    return heading_map


def _build_numbering_map(raw_bytes):
    num_fmt_map = {}
    try:
        root = etree.fromstring(raw_bytes)
    except Exception:
        return num_fmt_map
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    abstract_fmts = {}
    for abs_num in root.iter(f"{{{ns}}}abstractNum"):
        abs_id = abs_num.get(f"{{{ns}}}abstractNumId")
        if abs_id is None:
            continue
        levels = {}
        for lvl in abs_num.iter(f"{{{ns}}}lvl"):
            ilvl = lvl.get(f"{{{ns}}}ilvl")
            num_fmt_el = lvl.find(f"{{{ns}}}numFmt")
            if num_fmt_el is not None:
                levels[int(ilvl)] = num_fmt_el.get(f"{{{ns}}}val", "decimal")
        abstract_fmts[abs_id] = levels

    for num in root.iter(f"{{{ns}}}num"):
        num_id = num.get(f"{{{ns}}}numId")
        if num_id is None:
            continue
        abs_num_ref = num.find(f"{{{ns}}}abstractNumId")
        if abs_num_ref is None:
            continue
        abs_id = abs_num_ref.get(f"{{{ns}}}val")
        if abs_id in abstract_fmts:
            num_fmt_map[num_id] = abstract_fmts[abs_id]
    return num_fmt_map


def _build_list_style_map(raw_bytes, num_fmt_map):
    list_style_map = {}
    try:
        root = etree.fromstring(raw_bytes)
    except Exception:
        return list_style_map
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for style_el in root.iter(f"{{{ns}}}style"):
        style_id = style_el.get(f"{{{ns}}}styleId")
        num_pr = style_el.find(f"{{{ns}}}pPr/{{{ns}}}numPr")
        if num_pr is None:
            continue
        num_id_el = num_pr.find(f"{{{ns}}}numId")
        ilvl_el = num_pr.find(f"{{{ns}}}ilvl")
        if num_id_el is None:
            continue
        num_id = num_id_el.get(f"{{{ns}}}val")
        ilvl = ilvl_el.get(f"{{{ns}}}val") if ilvl_el is not None else "0"
        fmt = "decimal"
        if num_id in num_fmt_map and int(ilvl) in num_fmt_map[num_id]:
            fmt = num_fmt_map[num_id][int(ilvl)]
        list_style_map[style_id] = (num_id, int(ilvl), fmt)
    return list_style_map


def _run_to_span(r_elem):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    text_parts = []
    for child in r_elem.iter():
        tag = child.tag
        if tag == f"{{{ns}}}t":
            text_parts.append(child.text or "")
        elif tag == f"{{{ns}}}tab":
            text_parts.append("\t")
        elif tag == f"{{{ns}}}br":
            text_parts.append("\n")
    text = "".join(text_parts)

    rPr = r_elem.find(f"{{{ns}}}rPr")
    bold = italic = strike = code = False
    if rPr is not None:
        b_elem = rPr.find(f"{{{ns}}}b")
        if b_elem is not None:
            val = b_elem.get(f"{{{ns}}}val", "true")
            bold = val.lower() not in ("0", "false", "off")
        i_elem = rPr.find(f"{{{ns}}}i")
        if i_elem is not None:
            val = i_elem.get(f"{{{ns}}}val", "true")
            italic = val.lower() not in ("0", "false", "off")
        strike = rPr.find(f"{{{ns}}}strike") is not None
        rFonts = rPr.find(f"{{{ns}}}rFonts")
        if rFonts is not None:
            ascii_font = rFonts.get(f"{{{ns}}}ascii")
            if ascii_font and "Courier" in ascii_font:
                code = True

    return text, bold, italic, strike, code


def _merge_spans(spans):
    if not spans:
        return []
    merged = [spans[0]]
    for span in spans[1:]:
        prev = merged[-1]
        if (prev[1] == span[1] and prev[2] == span[2] and
                prev[3] == span[3] and prev[4] == span[4]):
            merged[-1] = (prev[0] + span[0], prev[1], prev[2], prev[3], prev[4])
        else:
            merged.append(span)
    return merged


def _extract_inline_md(p_elem):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    spans = []
    for r in p_elem.iter(f"{{{ns}}}r"):
        span = _run_to_span(r)
        if span[0]:
            spans.append(span)
    spans = _merge_spans(spans)

    parts = []
    for text, bold, italic, strike, code in spans:
        text = text.replace("**", "\\*\\*").replace("*", "\\*")
        text = text.replace("~", "\\~")
        if code:
            parts.append(f"`{text}`")
        elif bold and italic:
            parts.append(f"***{text}***")
        elif bold:
            parts.append(f"**{text}**")
        elif italic:
            parts.append(f"*{text}*")
        elif strike:
            parts.append(f"~~{text}~~")
        else:
            parts.append(text)

    return "".join(parts)


def _extract_plain_text(p_elem):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    parts = []
    for r in p_elem.iter(f"{{{ns}}}r"):
        for t in r.iter(f"{{{ns}}}t"):
            if t.text:
                parts.append(t.text)
        for tab in r.iter(f"{{{ns}}}tab"):
            parts.append("\t")
        for br in r.iter(f"{{{ns}}}br"):
            parts.append("\n")
    return "".join(parts)


def _table_to_md(tbl_elem):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    rows = []
    for tr in tbl_elem.iter(f"{{{ns}}}tr"):
        cells = []
        for tc in tr.iter(f"{{{ns}}}tc"):
            cell_text = ""
            for p in tc.iter(f"{{{ns}}}p"):
                cell_text += _extract_plain_text(p) + " "
            cells.append(cell_text.strip().replace("|", "\\|"))
        rows.append(cells)
    if not rows:
        return ""
    num_cols = max(len(r) for r in rows)
    lines = []
    lines.append("| " + " | ".join(rows[0][c] if c < len(rows[0]) else "" for c in range(num_cols)) + " |")
    lines.append("|" + "---|" * num_cols)
    for row in rows[1:]:
        lines.append("| " + " | ".join(row[c] if c < len(row) else "" for c in range(num_cols)) + " |")
    return "\n".join(lines)


def _paragraph_to_md(p_elem, heading_map, num_fmt_map, list_map, comments_map):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    pPr = p_elem.find(f"{{{ns}}}pPr")

    style_id = None
    if pPr is not None:
        pStyle = pPr.find(f"{{{ns}}}pStyle")
        if pStyle is not None:
            style_id = pStyle.get(f"{{{ns}}}val")

    # Check for code block (Courier New font on ALL runs, or Code Block style)
    is_code = False
    if style_id == "CodeBlock" or style_id == "Code":
        is_code = True
    else:
        code_runs = []
        total_runs = 0
        for r in p_elem.iter(f"{{{ns}}}r"):
            _, _, _, _, code = _run_to_span(r)
            t, _, _, _, _ = _run_to_span(r)
            if t.strip():
                total_runs += 1
                if code:
                    code_runs.append(True)
        if total_runs > 0 and len(code_runs) == total_runs:
            is_code = True

    # Check for numberPr (list item)
    num_id = None
    ilvl = 0
    num_fmt = None
    if pPr is not None:
        numPr = pPr.find(f"{{{ns}}}numPr")
        if numPr is not None:
            num_id_el = numPr.find(f"{{{ns}}}numId")
            if num_id_el is not None:
                num_id = num_id_el.get(f"{{{ns}}}val")
            ilvl_el = numPr.find(f"{{{ns}}}ilvl")
            if ilvl_el is not None:
                ilvl = int(ilvl_el.get(f"{{{ns}}}val", "0"))
            if num_id and num_id in num_fmt_map and ilvl in num_fmt_map[num_id]:
                num_fmt = num_fmt_map[num_id][ilvl]

    # Check style-based list info
    if num_id is None and style_id and style_id in list_map:
        sid, silvl, sfmt = list_map[style_id]
        num_id = sid
        ilvl = silvl
        num_fmt = sfmt

    # Check heading
    heading_level = None
    if style_id and style_id in heading_map:
        hl = heading_map[style_id]
        if 1 <= hl <= 6:
            heading_level = hl

    if heading_level is not None:
        text = _extract_inline_md(p_elem)
        text = re.sub(r"\*{1,3}", "", text)
        text = re.sub(r"`([^`]+)`", r"\1", text)
        return "#" * heading_level + " " + text + "\n", False

    if is_code:
        lines = []
        for r in p_elem.iter(f"{{{ns}}}r"):
            t, _, _, _, _ = _run_to_span(r)
            if t:
                lines.append(t)
        text = "\n".join(lines).rstrip("\n")
        return "```\n" + text + "\n```\n", False

    if num_fmt == "bullet":
        text = _extract_inline_md(p_elem)
        return "- " + text + "\n", False

    if num_fmt in ("decimal", "lowerLetter", "upperLetter", "lowerRoman", "upperRoman"):
        text = _extract_inline_md(p_elem)
        return "1. " + text + "\n", False

    text = _extract_inline_md(p_elem)
    if text.strip() == "":
        return "\n", False
    return text + "\n", False


def _load_comments(raw_bytes):
    comments = {}
    try:
        root = etree.fromstring(raw_bytes)
    except Exception:
        return comments
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for comm in root.iter(f"{{{ns}}}comment"):
        cid = comm.get(f"{{{ns}}}id")
        author = comm.get(f"{{{ns}}}author", "")
        date = comm.get(f"{{{ns}}}date", "")
        text_parts = []
        for t in comm.iter(f"{{{ns}}}t"):
            if t.text:
                text_parts.append(t.text)
        comments[cid] = {"author": author, "date": date, "text": "".join(text_parts)}
    return comments


def _build_comment_ranges(root):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body = root.find(f"{{{ns}}}body")
    if body is None:
        return {}
    comment_indices = {}
    para_idx = -1
    for child in body.iter():
        if child.tag == f"{{{ns}}}p":
            para_idx += 1
        if child.tag == f"{{{ns}}}commentRangeStart":
            cid = child.get(f"{{{ns}}}id")
            if cid:
                if para_idx not in comment_indices:
                    comment_indices[para_idx] = []
                comment_indices[para_idx].append(cid)
    return comment_indices


def _docx_to_markdown(docx_path):
    with open(docx_path, "rb") as f:
        raw_bytes = f.read()

    with zipfile.ZipFile(io.BytesIO(raw_bytes), "r") as z:
        names = z.namelist()
        doc_xml = z.read("word/document.xml")
        styles_xml = z.read("word/styles.xml") if "word/styles.xml" in names else b""
        num_xml = z.read("word/numbering.xml") if "word/numbering.xml" in names else b""
        comments_xml = z.read("word/comments.xml") if "word/comments.xml" in names else b""

    heading_map = _build_style_heading_map(styles_xml)
    num_fmt_map = _build_numbering_map(num_xml)
    list_map = _build_list_style_map(styles_xml, num_fmt_map)
    comment_data = _load_comments(comments_xml)

    root = etree.fromstring(doc_xml)
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body = root.find(f"{{{ns}}}body")

    comment_ranges = _build_comment_ranges(root)
    para_idx = -1
    output_lines = []

    # Process body children, handling tables
    for child in body:
        if child.tag == f"{{{ns}}}tbl":
            md = _table_to_md(child) + "\n"
            if md.strip():
                output_lines.append(md)
            continue

        if child.tag != f"{{{ns}}}p":
            continue

        pPr = child.find(f"{{{ns}}}pPr")
        # Skip empty paragraphs that are just structure
        has_text = False
        for r in child.iter(f"{{{ns}}}r"):
            for t in r.iter(f"{{{ns}}}t"):
                if t.text and t.text.strip():
                    has_text = True
                    break
        if not has_text:
            # Check if it's a list item or heading with no text
            style_id = None
            if pPr is not None:
                pStyle = pPr.find(f"{{{ns}}}pStyle")
                if pStyle is not None:
                    style_id = pStyle.get(f"{{{ns}}}val")
            if not style_id:
                continue

        para_idx += 1
        md, is_table = _paragraph_to_md(child, heading_map, num_fmt_map, list_map, comment_data)
        output_lines.append(md)

        if para_idx in comment_ranges:
            for cid in comment_ranges[para_idx]:
                if cid in comment_data:
                    c = comment_data[cid]
                    output_lines.append(f"<!-- COMMENT [{c['author']}]: {c['text']} -->\n")

    result = "".join(output_lines)
    result = re.sub(r"\n{3,}", "\n\n", result)
    result = result.rstrip("\n") + "\n"
    return result


def parse_docx_to_md(docx_path, md_path):
    result = _docx_to_markdown(docx_path)
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(result)


def _apply_replacements(docx_path):
    replacements = {"\u2014": "\u2013", "\u0451": "\u0435", "\u0401": "\u0415"}
    with zipfile.ZipFile(docx_path, "r") as zin:
        data = {name: zin.read(name) for name in zin.namelist()}
    for name in list(data.keys()):
        if name.endswith(".xml") or name.endswith(".rels"):
            try:
                root = etree.fromstring(data[name])
            except Exception:
                continue
            changed = False
            for el in root.iter():
                if el.text:
                    new_text = el.text
                    for old, new in replacements.items():
                        if old in new_text:
                            new_text = new_text.replace(old, new)
                            changed = True
                    if changed:
                        el.text = new_text
            if changed:
                data[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(docx_path, "w") as zout:
        for name, content in data.items():
            zout.writestr(name, content)


def _inject_comments(buf, out_path, comments):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"

    with zipfile.ZipFile(buf, "r") as z:
        data = {n: z.read(n) for n in z.namelist()}

    doc_root = etree.fromstring(data["word/document.xml"])
    body = doc_root.find(f"{{{ns}}}body")

    # Map para_idx to body child element index
    child_elements = list(body)
    para_to_child = {}
    child_idx = -1
    for ci, child in enumerate(child_elements):
        if child.tag == f"{{{ns}}}p":
            child_idx += 1
            para_to_child[child_idx] = ci
        elif child.tag == f"{{{ns}}}tbl":
            child_idx += 1
            para_to_child[child_idx] = ci

    # Parse existing comments.xml or create new
    comments_root = None
    next_id = 0
    if "word/comments.xml" in data:
        try:
            comments_root = etree.fromstring(data["word/comments.xml"])
            existing_ids = []
            for c in comments_root.iter(f"{{{ns}}}comment"):
                cid = c.get(f"{{{ns}}}id")
                if cid is not None:
                    existing_ids.append(int(cid))
            if existing_ids:
                next_id = max(existing_ids) + 1
        except Exception:
            comments_root = None

    if comments_root is None:
        comments_root = etree.fromstring(
            f'<w:comments xmlns:w="{ns}"></w:comments>'
        )

    now_utc = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    for para_idx, author, text in comments:
        if para_idx not in para_to_child:
            continue
        ci = para_to_child[para_idx]
        target_elem = child_elements[ci]
        # Ensure target is a paragraph
        if target_elem.tag != f"{{{ns}}}p":
            continue

        cid = str(next_id)
        next_id += 1

        # Create comment element
        comm = etree.SubElement(comments_root, f"{{{ns}}}comment")
        comm.set(f"{{{ns}}}id", cid)
        comm.set(f"{{{ns}}}author", author)
        comm.set(f"{{{ns}}}date", now_utc)
        cp = etree.SubElement(comm, f"{{{ns}}}p")
        cr = etree.SubElement(cp, f"{{{ns}}}r")
        ct = etree.SubElement(cr, f"{{{ns}}}t")
        ct.text = text

        # Get or create pPr
        pPr = target_elem.find(f"{{{ns}}}pPr")
        if pPr is None:
            pPr = etree.Element(f"{{{ns}}}pPr")
            target_elem.insert(0, pPr)

        # Insert commentRangeStart after pPr
        cs = etree.Element(f"{{{ns}}}commentRangeStart")
        cs.set(f"{{{ns}}}id", cid)
        pPr.addnext(cs)

        # Add commentRangeEnd and commentReference at the end
        ce = etree.Element(f"{{{ns}}}commentRangeEnd")
        ce.set(f"{{{ns}}}id", cid)
        target_elem.append(ce)

        ref_r = etree.SubElement(target_elem, f"{{{ns}}}r")
        ref_rPr = etree.SubElement(ref_r, f"{{{ns}}}rPr")
        ref_rStyle = etree.SubElement(ref_rPr, f"{{{ns}}}rStyle")
        ref_rStyle.set(f"{{{ns}}}val", "CommentReference")
        ref_cr = etree.SubElement(ref_r, f"{{{ns}}}commentReference")
        ref_cr.set(f"{{{ns}}}id", cid)

    data["word/document.xml"] = etree.tostring(doc_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    data["word/comments.xml"] = etree.tostring(comments_root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # Update [Content_Types].xml
    ct_root = etree.fromstring(data["[Content_Types].xml"])
    has_comment_ct = False
    for child in ct_root:
        if child.get("PartName") == "/word/comments.xml":
            has_comment_ct = True
            break
    if not has_comment_ct:
        override = etree.SubElement(ct_root, f"{{{ct_ns}}}Override")
        override.set("PartName", "/word/comments.xml")
        override.set("ContentType", "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml")
        data["[Content_Types].xml"] = etree.tostring(ct_root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # Update word/_rels/document.xml.rels
    rels_root = etree.fromstring(data["word/_rels/document.xml.rels"])
    has_comment_rel = False
    max_rel_id = 0
    for child in rels_root:
        rid = child.get("Id", "")
        if rid.startswith("rId"):
            try:
                max_rel_id = max(max_rel_id, int(rid[3:]))
            except ValueError:
                pass
        if child.get("Target") == "comments.xml":
            has_comment_rel = True
            break
    if not has_comment_rel:
        rel = etree.SubElement(rels_root, f"{{{rel_ns}}}Relationship")
        rel.set("Id", f"rId{max_rel_id + 1}")
        rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments")
        rel.set("Target", "comments.xml")
        data["word/_rels/document.xml.rels"] = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone=True)

    with zipfile.ZipFile(out_path, "w") as zout:
        for name, content in data.items():
            zout.writestr(name, content)


def _inject_numbering(docx_path):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    with zipfile.ZipFile(docx_path, "r") as z:
        data = {n: z.read(n) for n in z.namelist()}

    doc_root = etree.fromstring(data["word/document.xml"])
    body = doc_root.find(f"{{{ns}}}body")

    styles_root = etree.fromstring(data["word/styles.xml"])

    has_numbering = "word/numbering.xml" in data
    if has_numbering:
        num_root = etree.fromstring(data["word/numbering.xml"])
    else:
        num_root = etree.fromstring(
            f'<w:numbering xmlns:w="{ns}" xmlns:r="{r_ns}"></w:numbering>'
        )

    # Find max existing IDs
    max_abs_id = -1
    max_num_id = -1
    for abs_num in num_root.iter(f"{{{ns}}}abstractNum"):
        aid = abs_num.get(f"{{{ns}}}abstractNumId")
        if aid is not None:
            try:
                max_abs_id = max(max_abs_id, int(aid))
            except ValueError:
                pass
    for num_el in num_root.iter(f"{{{ns}}}num"):
        nid = num_el.get(f"{{{ns}}}numId")
        if nid is not None:
            try:
                max_num_id = max(max_num_id, int(nid))
            except ValueError:
                pass

    def next_abs_id():
        nonlocal max_abs_id
        max_abs_id += 1
        return str(max_abs_id)

    def next_num_id():
        nonlocal max_num_id
        max_num_id += 1
        return str(max_num_id)

    # === BULLET NUMBERING ===
    # Find existing bullet abstractNum
    bullet_abs_id = None
    for abs_num in num_root.iter(f"{{{ns}}}abstractNum"):
        for lvl in abs_num.iter(f"{{{ns}}}lvl"):
            nf = lvl.find(f"{{{ns}}}numFmt")
            if nf is not None and nf.get(f"{{{ns}}}val") == "bullet":
                bullet_abs_id = abs_num.get(f"{{{ns}}}abstractNumId")
                break
        if bullet_abs_id is not None:
            break

    if bullet_abs_id is None:
        bullet_abs_id = next_abs_id()
        b_abs = etree.SubElement(num_root, f"{{{ns}}}abstractNum")
        b_abs.set(f"{{{ns}}}abstractNumId", bullet_abs_id)
        for lvl in range(9):
            lvl_el = etree.SubElement(b_abs, f"{{{ns}}}lvl")
            lvl_el.set(f"{{{ns}}}ilvl", str(lvl))
            start_el = etree.SubElement(lvl_el, f"{{{ns}}}start")
            start_el.set(f"{{{ns}}}val", "1")
            nf_el = etree.SubElement(lvl_el, f"{{{ns}}}numFmt")
            nf_el.set(f"{{{ns}}}val", "bullet")
            lvl_text = etree.SubElement(lvl_el, f"{{{ns}}}lvlText")
            lvl_text.set(f"{{{ns}}}val", "\u2022")
            lvl_jc = etree.SubElement(lvl_el, f"{{{ns}}}lvlJc")
            lvl_jc.set(f"{{{ns}}}val", "left")

    # Find existing bullet num
    bullet_num_id = None
    for num_el in num_root.iter(f"{{{ns}}}num"):
        ref = num_el.find(f"{{{ns}}}abstractNumId")
        if ref is not None and ref.get(f"{{{ns}}}val") == bullet_abs_id:
            bullet_num_id = num_el.get(f"{{{ns}}}numId")
            break

    if bullet_num_id is None:
        bullet_num_id = next_num_id()
        b_num = etree.SubElement(num_root, f"{{{ns}}}num")
        b_num.set(f"{{{ns}}}numId", bullet_num_id)
        ref = etree.SubElement(b_num, f"{{{ns}}}abstractNumId")
        ref.set(f"{{{ns}}}val", bullet_abs_id)

    # Modify ListBullet and ListParagraph styles to use bullet numbering
    for style in styles_root.iter(f"{{{ns}}}style"):
        sid = style.get(f"{{{ns}}}styleId")
        if sid in ("ListBullet", "ListBullet2", "ListBullet3", "ListParagraph"):
            pPr = style.find(f"{{{ns}}}pPr")
            if pPr is None:
                pPr = etree.SubElement(style, f"{{{ns}}}pPr")
                style.append(pPr)
            numPr = pPr.find(f"{{{ns}}}numPr")
            if numPr is None:
                numPr = etree.SubElement(pPr, f"{{{ns}}}numPr")
                pPr.append(numPr)
            numId_el = numPr.find(f"{{{ns}}}numId")
            if numId_el is None:
                numId_el = etree.SubElement(numPr, f"{{{ns}}}numId")
                numPr.append(numId_el)
            numId_el.set(f"{{{ns}}}val", bullet_num_id)
            ilvl_el = numPr.find(f"{{{ns}}}ilvl")
            if ilvl_el is None:
                ilvl_el = etree.SubElement(numPr, f"{{{ns}}}ilvl")
                numPr.append(ilvl_el)
            ilvl_el.set(f"{{{ns}}}val", "0")

    # === HEADING AUTO-NUMBERING ===
    heading_abs_id = next_abs_id()
    h_abs = etree.SubElement(num_root, f"{{{ns}}}abstractNum")
    h_abs.set(f"{{{ns}}}abstractNumId", heading_abs_id)
    multiLevel_type = etree.SubElement(h_abs, f"{{{ns}}}multiLevelType")
    multiLevel_type.set(f"{{{ns}}}val", "hybridMultilevel")
    for lvl in range(9):
        lvl_el = etree.SubElement(h_abs, f"{{{ns}}}lvl")
        lvl_el.set(f"{{{ns}}}ilvl", str(lvl))
        start_el = etree.SubElement(lvl_el, f"{{{ns}}}start")
        start_el.set(f"{{{ns}}}val", "1")
        nf_el = etree.SubElement(lvl_el, f"{{{ns}}}numFmt")
        nf_el.set(f"{{{ns}}}val", "decimal")
        level_text_parts = [f"%{i+1}." for i in range(lvl + 1)]
        lvl_text = etree.SubElement(lvl_el, f"{{{ns}}}lvlText")
        lvl_text.set(f"{{{ns}}}val", "".join(level_text_parts))
        lvl_jc = etree.SubElement(lvl_el, f"{{{ns}}}lvlJc")
        lvl_jc.set(f"{{{ns}}}val", "left")

    heading_num_id = next_num_id()
    h_num = etree.SubElement(num_root, f"{{{ns}}}num")
    h_num.set(f"{{{ns}}}numId", heading_num_id)
    ref = etree.SubElement(h_num, f"{{{ns}}}abstractNumId")
    ref.set(f"{{{ns}}}val", heading_abs_id)

    # Assign heading numbering to Heading 1-9 styles
    for i in range(1, 10):
        style_id = f"Heading{i}"
        for style in styles_root.iter(f"{{{ns}}}style"):
            sid = style.get(f"{{{ns}}}styleId")
            if sid == style_id:
                pPr = style.find(f"{{{ns}}}pPr")
                if pPr is None:
                    pPr = etree.SubElement(style, f"{{{ns}}}pPr")
                    style.append(pPr)
                numPr = pPr.find(f"{{{ns}}}numPr")
                if numPr is None:
                    numPr = etree.SubElement(pPr, f"{{{ns}}}numPr")
                    pPr.append(numPr)
                numId_el = numPr.find(f"{{{ns}}}numId")
                if numId_el is None:
                    numId_el = etree.SubElement(numPr, f"{{{ns}}}numId")
                    numPr.append(numId_el)
                numId_el.set(f"{{{ns}}}val", heading_num_id)
                ilvl_el = numPr.find(f"{{{ns}}}ilvl")
                if ilvl_el is None:
                    ilvl_el = etree.SubElement(numPr, f"{{{ns}}}ilvl")
                    numPr.append(ilvl_el)
                ilvl_el.set(f"{{{ns}}}val", str(i - 1))
                break

    # === PER-GROUP NUMBERED LIST RESTART ===
    # Find list number numId from style
    list_num_style_id = None
    for style in styles_root.iter(f"{{{ns}}}style"):
        sid = style.get(f"{{{ns}}}styleId")
        if sid == "ListNumber":
            pPr = style.find(f"{{{ns}}}pPr")
            if pPr is not None:
                numPr = pPr.find(f"{{{ns}}}numPr")
                if numPr is not None:
                    nid_el = numPr.find(f"{{{ns}}}numId")
                    if nid_el is not None:
                        list_num_style_id = nid_el.get(f"{{{ns}}}val")
            break

    # Scan body for consecutive ListNumber paragraphs
    paragraphs = []
    for child in body:
        if child.tag == f"{{{ns}}}p":
            paragraphs.append(child)
        # Skip tables

    list_groups = []
    current_group = None
    for p in paragraphs:
        pPr = p.find(f"{{{ns}}}pPr")
        style_id_in_p = None
        if pPr is not None:
            pStyle = pPr.find(f"{{{ns}}}pStyle")
            if pStyle is not None:
                style_id_in_p = pStyle.get(f"{{{ns}}}val")
        if style_id_in_p == "ListNumber":
            if current_group is None:
                current_group = []
                list_groups.append(current_group)
            current_group.append(p)
        else:
            current_group = None

    for group in list_groups:
        g_num_id = next_num_id()
        g_num = etree.SubElement(num_root, f"{{{ns}}}num")
        g_num.set(f"{{{ns}}}numId", g_num_id)
        # Reference the same abstractNum as ListNumber
        if list_num_style_id:
            # Find the num that ListNumber style references
            ref_abs = None
            for num_el in num_root.iter(f"{{{ns}}}num"):
                nid = num_el.get(f"{{{ns}}}numId")
                if nid == list_num_style_id:
                    ref_abs_el = num_el.find(f"{{{ns}}}abstractNumId")
                    if ref_abs_el is not None:
                        ref_abs = ref_abs_el.get(f"{{{ns}}}val")
                    break
            if ref_abs is None:
                # Create a new abstractNum for list numbering
                list_abs_id = next_abs_id()
                l_abs = etree.SubElement(num_root, f"{{{ns}}}abstractNum")
                l_abs.set(f"{{{ns}}}abstractNumId", list_abs_id)
                for lvl in range(9):
                    lvl_el = etree.SubElement(l_abs, f"{{{ns}}}lvl")
                    lvl_el.set(f"{{{ns}}}ilvl", str(lvl))
                    start_el = etree.SubElement(lvl_el, f"{{{ns}}}start")
                    start_el.set(f"{{{ns}}}val", "1")
                    nf_el = etree.SubElement(lvl_el, f"{{{ns}}}numFmt")
                    nf_el.set(f"{{{ns}}}val", "decimal")
                    lvl_text = etree.SubElement(lvl_el, f"{{{ns}}}lvlText")
                    lvl_text.set(f"{{{ns}}}val", f"%{lvl+1}.")
                    lvl_jc = etree.SubElement(lvl_el, f"{{{ns}}}lvlJc")
                    lvl_jc.set(f"{{{ns}}}val", "left")
                ref_abs = list_abs_id
            ref = etree.SubElement(g_num, f"{{{ns}}}abstractNumId")
            ref.set(f"{{{ns}}}val", ref_abs)
        else:
            # Use the heading abstract num as fallback
            ref = etree.SubElement(g_num, f"{{{ns}}}abstractNumId")
            ref.set(f"{{{ns}}}val", heading_abs_id)

        # Add startOverride
        lvlOverride = etree.SubElement(g_num, f"{{{ns}}}lvlOverride")
        lvlOverride.set(f"{{{ns}}}ilvl", "0")
        startOverride = etree.SubElement(lvlOverride, f"{{{ns}}}startOverride")
        startOverride.set(f"{{{ns}}}val", "1")

        for p in group:
            pPr = p.find(f"{{{ns}}}pPr")
            if pPr is None:
                pPr = etree.Element(f"{{{ns}}}pPr")
                p.insert(0, pPr)
            # Remove existing numPr
            existing_numPr = pPr.find(f"{{{ns}}}numPr")
            if existing_numPr is not None:
                pPr.remove(existing_numPr)
            # Remove pStyle
            existing_pStyle = pPr.find(f"{{{ns}}}pStyle")
            if existing_pStyle is not None:
                pPr.remove(existing_pStyle)

            new_numPr = etree.SubElement(pPr, f"{{{ns}}}numPr")
            new_ilvl = etree.SubElement(new_numPr, f"{{{ns}}}ilvl")
            new_ilvl.set(f"{{{ns}}}val", "0")
            new_numId = etree.SubElement(new_numPr, f"{{{ns}}}numId")
            new_numId.set(f"{{{ns}}}val", g_num_id)
            pPr.append(new_numPr)

    data["word/document.xml"] = etree.tostring(doc_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    data["word/styles.xml"] = etree.tostring(styles_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    data["word/numbering.xml"] = etree.tostring(num_root, xml_declaration=True, encoding="UTF-8", standalone=True)

    with zipfile.ZipFile(docx_path, "w") as zout:
        for name, content in data.items():
            zout.writestr(name, content)


def main():
    if len(sys.argv) >= 2:
        raw = sys.argv[1]
        if raw.startswith("file://"):
            raw = raw[len("file://"):]
        src = os.path.expanduser(raw)
        src = os.path.abspath(src)

        if not os.path.isfile(src):
            print(f"Error: file not found: {src}", file=sys.stderr)
            sys.exit(1)

        in_ext = os.path.splitext(src)[1].lower()
        if in_ext in (".md", ".markdown"):
            out_ext = ".docx"
        elif in_ext == ".docx":
            out_ext = ".md"
        else:
            print(f"Error: unsupported file extension: {in_ext}", file=sys.stderr)
            sys.exit(1)

        out_path = os.path.splitext(src)[0] + out_ext
        print(f"Reading: {src}")
        print(f"Output: {out_path}")

        try:
            if out_ext == ".docx":
                parse_md_to_docx(src, out_path)
            else:
                parse_docx_to_md(src, out_path)
        except NotImplementedError as e:
            print(f"Error: {e}", file=sys.stderr)
            sys.exit(1)
        except Exception as e:
            print(f"Error: {e}", file=sys.stderr)
            sys.exit(1)

        print("Done.")
        sys.exit(0)
    else:
        if USE_GUI:
            root = TkinterDnD.Tk()
            from tkinterdnd2 import DND_FILES
            app = App(root)
            root.mainloop()
        else:
            print("Error: GUI requires tkinterdnd2. Install with: pip install tkinterdnd2", file=sys.stderr)
            sys.exit(1)


class App:
    def __init__(self, root):
        self.root = root
        self.file_path = None
        self.queue = queue.Queue()

        root.title("MD \u2194 DOCX")
        root.geometry("440x280")
        root.resizable(False, False)

        self.btn_select = tk.Button(root, text="Select File\u2026", font=("", 14), relief=tk.GROOVE, command=self.browse)
        self.btn_select.pack(pady=(15, 5))

        self.lbl_path = tk.Label(root, text="No file selected", fg="grey", wraplength=380)
        self.lbl_path.pack(pady=2)

        self.lbl_dir = tk.Label(root, text="", fg="#0563C1")
        self.lbl_dir.pack(pady=2)

        self.lbl_status = tk.Label(root, text="Ready", fg="grey")
        self.lbl_status.pack(pady=2)

        self.btn_convert = tk.Button(root, text="Convert", font=("", 14), command=self.convert)
        self.btn_convert.pack(fill=tk.X, padx=20, pady=5)

        self.progress = tk.ttk.Progressbar(root, mode="indeterminate", length=400)
        self.progress.pack(pady=5)
        self.progress.pack_forget()

        root.drop_target_register(DND_FILES)
        root.dnd_bind("<<Drop>>", self._on_drop)

        self.root.after(80, self._poll_queue)

    def _parse_dnd_paths(self, raw):
        result = []
        current = ""
        in_braces = False
        for ch in raw:
            if ch == "{":
                in_braces = True
            elif ch == "}":
                in_braces = False
                result.append(current)
                current = ""
            elif ch == " " and not in_braces:
                if current:
                    result.append(current)
                    current = ""
            else:
                current += ch
        if current:
            result.append(current)
        return result

    def _on_drop(self, event):
        paths = self._parse_dnd_paths(event.data)
        if paths:
            self._select_file(paths[0])

    def _select_file(self, path):
        if path.startswith("file://"):
            path = path[len("file://"):]
        ext = os.path.splitext(path)[1].lower()
        if ext not in (".md", ".markdown", ".docx"):
            self.lbl_path.config(text="Unsupported file type", fg="red")
            return
        self.file_path = path
        fname = os.path.basename(path)
        self.lbl_path.config(text=fname, fg="black")
        if ext in (".md", ".markdown"):
            self.lbl_dir.config(text="MD \u2192 DOCX", fg="#0563C1")
        else:
            self.lbl_dir.config(text="DOCX \u2192 MD", fg="#0563C1")
        self.lbl_status.config(text="Ready", fg="grey")

    def browse(self):
        path = filedialog.askopenfilename(
            title="Select file",
            filetypes=[
                ("Supported files", "*.md *.markdown *.docx"),
                ("Markdown", "*.md *.markdown"),
                ("Word Document", "*.docx"),
                ("All files", "*.*"),
            ],
        )
        if path:
            self._select_file(path)

    def convert(self):
        if self.file_path is None:
            return
        ext = os.path.splitext(self.file_path)[1].lower()
        if ext in (".md", ".markdown"):
            out_path = os.path.splitext(self.file_path)[0] + ".docx"
        else:
            out_path = os.path.splitext(self.file_path)[0] + ".md"

        if os.path.exists(out_path):
            if not messagebox.askyesno("Overwrite?", f"Output file exists:\n{out_path}\nOverwrite?"):
                return

        self.btn_convert.config(state=tk.DISABLED)
        self.lbl_status.config(text="Converting\u2026", fg="grey")
        self.progress.pack(pady=5)
        self.progress.start()

        def worker():
            try:
                if ext in (".md", ".markdown"):
                    parse_md_to_docx(self.file_path, out_path)
                else:
                    parse_docx_to_md(self.file_path, out_path)
                self.queue.put(("done", out_path))
            except Exception as e:
                self.queue.put(("error", str(e)))

        t = threading.Thread(target=worker, daemon=True)
        t.start()

    def _poll_queue(self):
        try:
            msg = self.queue.get_nowait()
            kind, data = msg
            if kind == "done":
                self._on_done(data)
            elif kind == "error":
                self._on_error(data)
        except queue.Empty:
            pass
        finally:
            self.root.after(80, self._poll_queue)

    def _on_done(self, out_path):
        self.progress.stop()
        self.progress.pack_forget()
        self.btn_convert.config(state=tk.NORMAL)
        self.lbl_status.config(text=f"Saved: {os.path.basename(out_path)}", fg="green")

    def _on_error(self, msg):
        self.progress.stop()
        self.progress.pack_forget()
        self.btn_convert.config(state=tk.NORMAL)
        self.lbl_status.config(text=f"Error: {msg}", fg="red")


if __name__ == "__main__":
    main()
